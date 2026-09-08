#!/usr/bin/env python3
"""docx → Markdown с картинками, без pandoc (только python-docx).

  materials/docx2md.py                  in/ → out/ рядом со скриптом
  materials/docx2md.py --in A --out B   другие папки

Для каждого X.docx (рекурсивно, структура папок сохраняется):
  out/<путь>/X.md                 текст
  out/<путь>/X_images/img-001.png картинки, на них ссылается md

Понимает заголовки (в том числе русские стили «Заголовок N»), жирный и курсив,
маркированные и нумерованные списки с вложенностью, таблицы, ссылки, картинки.
EMF/WMF сохраняются как есть, без конвертации. Текстовые врезки (text box)
и сноски не извлекаются.

Зависимость: pip install python-docx
"""
from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.table import Table
    from docx.text.hyperlink import Hyperlink
    from docx.text.paragraph import Paragraph
    from docx.text.run import Run
except ImportError:
    sys.exit("нужен python-docx:  pip install python-docx")

HERE = Path(__file__).resolve().parent
VML_IMAGEDATA = "{urn:schemas-microsoft-com:vml}imagedata"
R_ID = qn("r:id")
R_EMBED = qn("r:embed")


# --- вспомогательное ------------------------------------------------------------

def safe_stem(stem: str) -> str:
    s = re.sub(r"[\\/:*?\"<>|]+", "", stem).strip()
    return re.sub(r"\s+", "_", s) or "doc"


class Numbering:
    """Формат списка (bullet / decimal) по numId и уровню, из numbering.xml."""

    def __init__(self, doc):
        self.root = None
        try:
            self.root = doc.part.numbering_part.element
        except Exception:  # noqa: BLE001  (в документе нет списков)
            pass
        self.cache: dict[tuple[str, str], str] = {}

    def fmt(self, num_id: str, ilvl: str) -> str:
        key = (num_id, ilvl)
        if key in self.cache:
            return self.cache[key]
        result = "bullet"
        if self.root is not None:
            abstract = None
            for num in self.root.findall(qn("w:num")):
                if num.get(qn("w:numId")) == num_id:
                    el = num.find(qn("w:abstractNumId"))
                    abstract = el.get(qn("w:val")) if el is not None else None
                    break
            if abstract is not None:
                for an in self.root.findall(qn("w:abstractNum")):
                    if an.get(qn("w:abstractNumId")) != abstract:
                        continue
                    for lvl in an.findall(qn("w:lvl")):
                        if lvl.get(qn("w:ilvl")) == ilvl:
                            nf = lvl.find(qn("w:numFmt"))
                            if nf is not None and nf.get(qn("w:val")) not in ("bullet", None):
                                result = "decimal"
                    break
        self.cache[key] = result
        return result


class Converter:
    def __init__(self, path: Path, img_dir: Path, img_rel: str):
        self.doc = Document(str(path))
        self.img_dir = img_dir
        self.img_rel = img_rel
        self.numbering = Numbering(self.doc)
        self.images: dict[str, str] = {}  # rId → имя файла
        self.counters: dict[str, dict[int, int]] = {}
        self.stats = {"paragraphs": 0, "images": 0, "tables": 0, "unconverted_images": 0}
        self.has_title = any(
            (p.style is not None and (p.style.name or "").lower() in ("title", "название")) and p.text.strip()
            for p in self.doc.paragraphs
        )

    # --- картинки ---

    def _save_image(self, rid: str) -> str | None:
        if rid in self.images:
            return self.images[rid]
        try:
            part = self.doc.part.related_parts[rid]
        except KeyError:
            return None
        ext = Path(str(part.partname)).suffix.lower() or ".bin"
        name = f"img-{len(self.images) + 1:03d}{ext}"
        self.img_dir.mkdir(parents=True, exist_ok=True)
        (self.img_dir / name).write_bytes(part.blob)
        self.images[rid] = name
        self.stats["images"] += 1
        if ext in (".emf", ".wmf"):
            self.stats["unconverted_images"] += 1
        return name

    def _images_in(self, element) -> list[str]:
        rids = [b.get(R_EMBED) for b in element.iter(qn("a:blip"))]
        rids += [i.get(R_ID) for i in element.iter(VML_IMAGEDATA)]
        out = []
        for rid in rids:
            if rid:
                name = self._save_image(rid)
                if name:
                    out.append(f"![]({self.img_rel}/{name})")
        return out

    # --- текст ---

    @staticmethod
    def _run_flags(run: Run) -> tuple[bool, bool]:
        bold, italic = run.bold, run.italic
        try:
            font = run.style.font if run.style is not None else None
            if bold is None and font is not None:
                bold = font.bold
            if italic is None and font is not None:
                italic = font.italic
        except Exception:  # noqa: BLE001
            pass
        return bool(bold), bool(italic)

    def _inline(self, paragraph: Paragraph) -> str:
        segs: list[tuple[str, bool, bool]] = []
        pictures: list[str] = []
        for item in paragraph.iter_inner_content():
            if isinstance(item, Hyperlink):
                text = item.text.strip()
                if text:
                    segs.append((f"[{text}]({item.address})" if item.address else text, False, False))
                for run in item.runs:
                    pictures += self._images_in(run._r)
            elif isinstance(item, Run):
                pictures += self._images_in(item._r)
                text = item.text.replace("\t", " ")
                if text:
                    b, i = self._run_flags(item)
                    segs.append((text, b, i))
        # склеиваем соседние куски с одинаковым форматированием
        merged: list[list] = []
        for text, b, i in segs:
            if merged and merged[-1][1] == b and merged[-1][2] == i:
                merged[-1][0] += text
            else:
                merged.append([text, b, i])
        text = self._render(merged).replace("\n", "  \n")
        text = re.sub(r"[ ]{2,}", " ", text).strip()
        if pictures:
            text = (text + "\n\n" if text else "") + "\n\n".join(pictures)
        return text

    @staticmethod
    def _wrap(text: str, marker: str) -> str:
        """Обернуть в маркер, оставив пробелы по краям снаружи; пустое не оборачиваем."""
        core = text.strip()
        if not core:
            return text
        lead = text[: len(text) - len(text.lstrip())]
        trail = text[len(text.rstrip()):]
        return f"{lead}{marker}{core}{marker}{trail}"

    @staticmethod
    def _render(segs: list[list]) -> str:
        """Маркеры как стек: закрываем и переоткрываем так, чтобы вложенность была правильной
        (*a **b** c*, а не *a* ***b**** c*). Пробелы по краям кусков выносим за маркеры."""
        out: list[str] = []
        stack: list[str] = []
        pending_ws = ""

        def wants(k: int) -> list[str]:
            return (["**"] if segs[k][1] else []) + (["*"] if segs[k][2] else [])

        def end_of(marker: str, start: int) -> int:
            """Индекс куска, где стиль заканчивается; пробельные куски не считаются."""
            k = start
            while k < len(segs) and (not segs[k][0].strip() or marker in wants(k)):
                k += 1
            return k

        for idx, (text, b, i) in enumerate(segs):
            if not text.strip():
                pending_ws += text
                continue
            want = wants(idx)
            extra = [m for m in stack if m not in want]
            reopen: list[str] = []
            if extra:
                idx = min(stack.index(m) for m in extra)
                closing = stack[idx:]
                out.extend(reversed(closing))
                stack = stack[:idx]
                reopen = [m for m in closing if m in want]
            out.append(pending_ws)
            pending_ws = ""
            lead = text[: len(text) - len(text.lstrip())]
            core = text.strip()
            trail = text[len(text.rstrip()):]
            to_open = reopen + [m for m in want if m not in stack and m not in reopen]
            # тот, что живёт дольше, открываем снаружи: *a **b** c*, а не **...****
            to_open.sort(key=lambda m: end_of(m, idx), reverse=True)
            out.append(lead)
            for m in to_open:
                out.append(m)
                stack.append(m)
            out.append(core)
            pending_ws = trail
        out.extend(reversed(stack))
        out.append(pending_ws)
        return "".join(out)

    def _heading_level(self, p: Paragraph) -> int:
        name = ""
        try:
            name = p.style.name or ""
        except Exception:  # noqa: BLE001
            pass
        # Если в документе есть стиль Title (название главы), Heading N уходит на уровень ниже.
        shift = 1 if self.has_title else 0
        m = re.match(r"(?:heading|заголовок)\s*(\d)", name, re.I)
        if m:
            return int(m.group(1)) + shift
        if name.lower() in ("title", "название"):
            return 1
        ppr = p._p.pPr
        if ppr is not None:
            lvl = ppr.find(qn("w:outlineLvl"))
            if lvl is not None:
                return int(lvl.get(qn("w:val"))) + 1 + shift
        return 0

    def _list_prefix(self, p: Paragraph) -> str | None:
        ppr = p._p.pPr
        num_pr = ppr.find(qn("w:numPr")) if ppr is not None else None
        style = ""
        try:
            style = (p.style.name or "").lower()
        except Exception:  # noqa: BLE001
            pass
        if num_pr is None:
            if style.startswith("list") or style.startswith("список"):
                return "- "
            return None
        num_id_el = num_pr.find(qn("w:numId"))
        ilvl_el = num_pr.find(qn("w:ilvl"))
        num_id = num_id_el.get(qn("w:val")) if num_id_el is not None else "0"
        ilvl = ilvl_el.get(qn("w:val")) if ilvl_el is not None else "0"
        level = int(ilvl)
        indent = "  " * level
        if self.numbering.fmt(num_id, ilvl) == "bullet":
            return indent + "- "
        counters = self.counters.setdefault(num_id, {})
        counters[level] = counters.get(level, 0) + 1
        for deeper in [k for k in counters if k > level]:
            counters.pop(deeper)
        return f"{indent}{counters[level]}. "

    def paragraph(self, p: Paragraph) -> str:
        text = self._inline(p)
        if not text:
            return ""
        self.stats["paragraphs"] += 1
        level = self._heading_level(p)
        if level:
            plain = re.sub(r"(\*{1,3})(\S.*?\S|\S)\1", r"\2", text)  # заголовок без жирного и курсива
            return "#" * min(level, 6) + " " + plain
        prefix = self._list_prefix(p)
        if prefix is not None:
            return prefix + text
        return text

    def table(self, t: Table) -> str:
        self.stats["tables"] += 1
        rows = []
        # Объединённые ячейки python-docx повторяет (тот же <w:tc>). Держим ссылки на элементы,
        # иначе lxml пересоздаёт прокси и сравнение по id() врёт.
        seen_tc: list = []
        for row in t.rows:
            cells = []
            prev_tc = None
            for cell in row.cells:
                tc = cell._tc
                if prev_tc is not None and tc is prev_tc:
                    continue  # горизонтальное объединение
                prev_tc = tc
                if any(tc is s for s in seen_tc):
                    cells.append("")  # вертикальное объединение
                    continue
                seen_tc.append(tc)
                lines = [self._inline(par) for par in cell.paragraphs]
                cells.append("<br>".join(x for x in lines if x).replace("|", "\\|").replace("\n", " "))
            rows.append(cells)
        rows = [r for r in rows if any(c.strip() for c in r)]  # строки-продолжения объединений
        captions = []
        while rows and sum(1 for c in rows[0] if c.strip()) == 1 and len(rows) > 1:
            captions.append(next(c for c in rows[0] if c.strip()))  # строка из одной ячейки = подпись
            rows = rows[1:]
        if not rows:
            return "\n\n".join(captions)
        width = max(len(r) for r in rows)
        rows = [r + [""] * (width - len(r)) for r in rows]
        out = [f"**{re.sub(r'^\*+|\*+$', '', c)}**" for c in captions]
        out += ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
        out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
        return "\n".join(out[:len(captions)]) + ("\n\n" if captions else "") + "\n".join(out[len(captions):])

    def convert(self) -> str:
        blocks = []
        body = self.doc.element.body
        for child in body.iterchildren():
            if child.tag == qn("w:p"):
                blocks.append(self.paragraph(Paragraph(child, self.doc)))
            elif child.tag == qn("w:tbl"):
                blocks.append(self.table(Table(child, self.doc)))
        text = "\n\n".join(b for b in blocks if b)
        # соседние пункты списка без пустой строки между ними
        text = re.sub(r"(\n(?:\s*(?:- |\d+\. ))[^\n]*)\n\n(?=\s*(?:- |\d+\. ))", r"\1\n", text)
        return text.strip() + "\n"


def convert_file(src: Path, dst_md: Path) -> dict:
    img_dir = dst_md.with_name(dst_md.stem + "_images")
    conv = Converter(src, img_dir, img_dir.name)
    md = conv.convert()
    dst_md.parent.mkdir(parents=True, exist_ok=True)
    dst_md.write_text(md, encoding="utf-8")
    return conv.stats


def main() -> int:
    ap = argparse.ArgumentParser(description="docx → md с картинками")
    ap.add_argument("--in", dest="src", default=str(HERE / "in"))
    ap.add_argument("--out", dest="dst", default=str(HERE / "out"))
    args = ap.parse_args()
    src, dst = Path(args.src), Path(args.dst)
    files = sorted(p for p in src.rglob("*.docx") if not p.name.startswith("~$"))
    if not files:
        sys.exit(f"в {src} нет .docx")
    total = {"paragraphs": 0, "images": 0, "tables": 0, "unconverted_images": 0}
    for f in files:
        rel = f.relative_to(src)
        out_md = dst / rel.parent / (safe_stem(f.stem) + ".md")
        try:
            stats = convert_file(f, out_md)
        except Exception as e:  # noqa: BLE001
            print(f"ОШИБКА {rel}: {e}")
            continue
        for k in total:
            total[k] += stats[k]
        print(f"{rel} → {out_md.relative_to(dst)}: абзацев {stats['paragraphs']}, "
              f"картинок {stats['images']}, таблиц {stats['tables']}"
              + (f", emf/wmf без конвертации {stats['unconverted_images']}" if stats["unconverted_images"] else ""))
    print(f"\nИтого: файлов {len(files)}, абзацев {total['paragraphs']}, картинок {total['images']}, таблиц {total['tables']}")
    if total["unconverted_images"]:
        print(f"emf/wmf сохранены как есть: {total['unconverted_images']} (для md-просмотрщиков их надо конвертировать в png)")
    return 0


if __name__ == "__main__":
    sys.exit(main())
